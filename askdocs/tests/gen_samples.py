"""Generates the fictional sample documents (Lumen Trail Outdoor Co. is a made-up company)."""
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
import docx

import os
os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "sample-docs"))

def pdf(path, pages):
    c = canvas.Canvas(path, pagesize=A4)
    for title, lines in pages:
        y = 800
        c.setFont("Helvetica-Bold", 15); c.drawString(60, y, title); y -= 30
        c.setFont("Helvetica", 11)
        for ln in lines:
            c.drawString(60, y, ln); y -= 18
        c.showPage()
    c.save()

pdf("warranty-handbook.pdf", [
 ("Lumen Trail Warranty Handbook", [
  "Lumen Trail Outdoor Co. is a fictional company used for this demo.",
  "All tents and backpacks carry a 3-year warranty against manufacturing defects.",
  "Headlamps and electronics carry a 2-year warranty.",
  "The warranty does not cover wear and tear, misuse, or damage from open flames."]),
 ("Making a warranty claim", [
  "To make a claim, email warranty@lumentrail.example with your order number and photos.",
  "Claims are reviewed within 5 business days.",
  "If the claim is approved we repair the item, or replace it if a repair is not possible.",
  "Return shipping for approved claims is paid by Lumen Trail."]),
])

d = docx.Document()
d.add_heading("Employee Onboarding Guide", 0)
d.add_heading("First day", 1)
d.add_paragraph("New team members receive a laptop and a badge from the operations desk on their first morning.")
d.add_paragraph("Your onboarding buddy will walk you through the warehouse safety tour before lunch.")
d.add_heading("Expenses", 1)
d.add_paragraph("Submit expense claims in the finance portal within 30 days of the purchase. Meals during business travel are reimbursed up to 40 USD per day.")
d.add_heading("Remote work", 1)
d.add_paragraph("Office staff may work remotely up to 2 days per week after their probation period of 3 months.")
t = d.add_table(rows=3, cols=2)
for r, (a, b) in enumerate([("Leave type", "Days per year"), ("Paid vacation", "25"), ("Sick leave", "10")]):
    t.cell(r, 0).text = a; t.cell(r, 1).text = b
d.save("employee-onboarding.docx")
