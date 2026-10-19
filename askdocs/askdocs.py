#!/usr/bin/env python3
"""
askdocs - chat with YOUR documents, with citations.

Load PDF, DOCX, Markdown, TXT and HTML files, build a local search index,
then ask questions. Every answer lists the file and page/section it came from.

Answer modes (pick one with --provider or the ASKDOCS_PROVIDER variable):
  none    Extractive. Returns the best matching passages. 100 % offline, no AI, no key.
  ollama  A local AI model through Ollama. 100 % local: nothing leaves your machine.
  openai  Any OpenAI-compatible API, with YOUR own API key (OpenAI, Mistral, Groq,
          OpenRouter, Azure OpenAI, a self-hosted vLLM...). Your question and the
          retrieved excerpts are sent to that provider.

Usage:
  python askdocs.py index ./my-docs
  python askdocs.py ask "What is the refund window?"
  python askdocs.py chat --provider ollama
  python askdocs.py serve            # local web chat on http://127.0.0.1:8765
  python askdocs.py export-web ../data/docs.json   # feed the browser widget

License: GPL-2.0-or-later. Contact: LaborX (https://laborx.com/freelancers/users/id475124)
"""
from __future__ import annotations

import argparse
import html
import logging
import json
import math
import os
import re
import sys
import time
import urllib.error
import urllib.request
from collections import Counter
from dataclasses import dataclass, asdict
from pathlib import Path

VERSION = "1.0.0"
SUPPORTED = {".pdf", ".docx", ".md", ".markdown", ".txt", ".html", ".htm"}
DEFAULT_INDEX_DIR = ".askdocs"
CHUNK_WORDS = 160
OVERLAP_WORDS = 30

# --------------------------------------------------------------------------- text extraction


@dataclass
class Section:
    file: str
    loc: str      # "page 3", "section: Refunds", "lines 10-24"
    title: str
    text: str


def _clean(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u00ad", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(path: Path, rel: str) -> list[Section]:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("PDF support needs pypdf: pip install -r requirements.txt")
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    reader = PdfReader(str(path))
    out = []
    for i, page in enumerate(reader.pages, start=1):
        txt = _clean(page.extract_text() or "")
        if txt:
            out.append(Section(rel, f"page {i}", path.stem, txt))
    return out


def extract_docx(path: Path, rel: str) -> list[Section]:
    try:
        import docx  # python-docx
    except ImportError:
        raise RuntimeError("DOCX support needs python-docx: pip install -r requirements.txt")
    d = docx.Document(str(path))
    out, heading, buf = [], path.stem, []

    def flush():
        txt = _clean("\n".join(buf))
        if txt:
            out.append(Section(rel, f"section: {heading}", heading, txt))

    for p in d.paragraphs:
        style = (p.style.name or "").lower() if p.style is not None else ""
        t = p.text.strip()
        if not t:
            continue
        if style.startswith("heading") or style == "title":
            flush()
            heading, buf = t, []
        else:
            buf.append(t)
    flush()
    for ti, table in enumerate(d.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        txt = _clean("\n".join(rows))
        if txt:
            out.append(Section(rel, f"table {ti}", f"{path.stem} table {ti}", txt))
    return out


def _split_markdown(text: str, rel: str, default_title: str) -> list[Section]:
    out, heading, buf = [], default_title, []

    def flush():
        txt = _clean("\n".join(buf))
        if txt:
            out.append(Section(rel, f"section: {heading}", heading, txt))

    for line in text.splitlines():
        m = re.match(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$", line)
        if m:
            flush()
            heading, buf = m.group(1).strip() or default_title, []
        else:
            buf.append(re.sub(r"[*_`]{1,3}", "", line))
    flush()
    return out


def extract_md(path: Path, rel: str) -> list[Section]:
    return _split_markdown(path.read_text(encoding="utf-8", errors="replace"), rel, path.stem)


def extract_txt(path: Path, rel: str) -> list[Section]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    out, start, buf = [], None, []
    for n, line in enumerate(lines + [""], start=1):
        if line.strip():
            if start is None:
                start = n
            buf.append(line)
        elif buf:
            out.append(Section(rel, f"lines {start}-{n - 1}", path.stem, _clean("\n".join(buf))))
            start, buf = None, []
    # merge tiny paragraphs so chunks are not 1 line each
    merged: list[Section] = []
    for s in out:
        if merged and len(merged[-1].text.split()) < 25:
            prev = merged[-1]
            a = prev.loc.split()[1].split("-")[0]
            b = s.loc.split()[1].split("-")[1]
            merged[-1] = Section(rel, f"lines {a}-{b}", prev.title, prev.text + "\n" + s.text)
        else:
            merged.append(s)
    return merged


def extract_html(path: Path, rel: str) -> list[Section]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        BeautifulSoup = None
    if BeautifulSoup is None:  # stdlib fallback: strip tags, keep headings as markdown
        raw = re.sub(r"(?is)<(script|style|noscript)[^>]*>.*?</\1>", " ", raw)
        raw = re.sub(r"(?is)<h([1-3])[^>]*>(.*?)</h\1>", lambda m: "\n# " + re.sub("<[^>]+>", "", m.group(2)) + "\n", raw)
        raw = re.sub(r"(?is)<(br|/p|/li|/div|/tr)[^>]*>", "\n", raw)
        text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
        return _split_markdown(text, rel, path.stem)
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "form"]):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else path.stem
    lines = []
    body = soup.body or soup
    for el in body.find_all(["h1", "h2", "h3", "h4", "p", "li", "td", "th", "pre", "blockquote", "dt", "dd"]):
        t = el.get_text(" ", strip=True)
        if not t:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            lines.append("# " + t)
        elif el.find_parent(["p", "li", "td", "th", "blockquote", "dd"]) is None or el.name in ("li",):
            lines.append(t)
    return _split_markdown("\n".join(lines), rel, title)


EXTRACTORS = {
    ".pdf": extract_pdf, ".docx": extract_docx, ".md": extract_md, ".markdown": extract_md,
    ".txt": extract_txt, ".html": extract_html, ".htm": extract_html,
}

# --------------------------------------------------------------------------- chunking + BM25

STOP = set("""a an the is are was were be been being do does did how what when where why which who whom
your you my our i me we they them it its of in on at to for from with by about into over after before
and or but if than then so as this that these those can could would should will just only also very
more most any all each every there here has have had not no yes may might must shall please tell""".split())


def tokenize(text: str) -> list[str]:
    toks = re.findall(r"[^\W_]+", text.lower(), flags=re.UNICODE)
    return [_stem(t) for t in toks]


def _stem(t: str) -> str:
    if len(t) <= 3 or t.isdigit():
        return t
    for suf, rep in (("ies", "y"), ("ing", ""), ("ed", ""), ("es", ""), ("s", "")):
        if t.endswith(suf) and len(t) - len(suf) >= 3 and not t.endswith("ss"):
            return t[: -len(suf)] + rep
    return t


def content_tokens(text: str) -> list[str]:
    return [t for t in tokenize(text) if t not in STOP and len(t) > 1]


def chunk_sections(sections: list[Section]) -> list[dict]:
    chunks = []
    for s in sections:
        words = s.text.split()
        if not words:
            continue
        step = CHUNK_WORDS - OVERLAP_WORDS
        starts = range(0, max(len(words) - OVERLAP_WORDS, 1), step) if len(words) > CHUNK_WORDS else [0]
        for part, st in enumerate(starts, start=1):
            piece = " ".join(words[st: st + CHUNK_WORDS])
            cid = f"{s.file}#{s.loc}" + (f"#{part}" if len(starts) > 1 else "")
            chunks.append({"id": cid, "file": s.file, "loc": s.loc, "title": s.title, "text": piece})
    return chunks


class BM25:
    def __init__(self, chunks: list[dict], k1: float = 1.5, b: float = 0.75):
        self.chunks, self.k1, self.b = chunks, k1, b
        self.tfs, self.lens = [], []
        df: Counter = Counter()
        for c in chunks:
            title_toks = content_tokens(c["title"] + " " + c["loc"])
            toks = content_tokens(c["text"]) + title_toks * 2  # small title boost
            tf = Counter(toks)
            self.tfs.append(tf)
            self.lens.append(max(len(toks), 1))
            df.update(tf.keys())
        self.df, self.N = df, len(chunks)
        self.avgdl = sum(self.lens) / max(self.N, 1)

    def search(self, query: str, k: int = 4) -> list[tuple[float, dict]]:
        q = list(dict.fromkeys(content_tokens(query)))
        if not q:
            return []
        need = math.ceil(len(q) / 2)  # at least half of the question's words must match
        scores = []
        for i, tf in enumerate(self.tfs):
            s, seen = 0.0, 0
            for t in q:
                f = tf.get(t)
                if not f:
                    continue
                seen += 1
                n = self.df[t]
                idf = math.log(1 + (self.N - n + 0.5) / (n + 0.5))
                s += idf * f * (self.k1 + 1) / (f + self.k1 * (1 - self.b + self.b * self.lens[i] / self.avgdl))
            if s > 0 and seen >= need:
                scores.append((s, self.chunks[i]))
        scores.sort(key=lambda x: -x[0])
        return scores[:k]

# --------------------------------------------------------------------------- index I/O


def build_index(paths: list[str], out_dir: str) -> dict:
    files, skipped, sections = [], [], []
    roots = [Path(p).expanduser().resolve() for p in paths]
    for root in roots:
        if not root.exists():
            raise SystemExit(f"Not found: {root}")
        candidates = [root] if root.is_file() else sorted(p for p in root.rglob("*") if p.is_file())
        base = root.parent if root.is_file() else root
        for p in candidates:
            if any(part.startswith(".") for part in p.relative_to(base).parts):
                continue
            ext = p.suffix.lower()
            if ext not in SUPPORTED:
                continue
            rel = p.relative_to(base).as_posix()
            try:
                secs = EXTRACTORS[ext](p, rel)
            except Exception as e:  # keep going, report at the end
                skipped.append({"file": rel, "reason": str(e)})
                continue
            if not secs:
                reason = "no text found" + (" (scanned PDF? OCR is not included)" if ext == ".pdf" else "")
                skipped.append({"file": rel, "reason": reason})
                continue
            files.append(rel)
            sections.extend(secs)
    chunks = chunk_sections(sections)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    data = {"version": VERSION, "created": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
            "sources": [str(r) for r in roots], "files": files, "skipped": skipped, "chunks": chunks}
    (out / "index.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    return data


def load_index(index_dir: str) -> dict:
    p = Path(index_dir) / "index.json"
    if not p.exists():
        raise SystemExit(f"No index at {p}. Run first:  python askdocs.py index <your-docs-folder>")
    return json.loads(p.read_text(encoding="utf-8"))

# --------------------------------------------------------------------------- answering

SYSTEM_PROMPT = (
    "You answer questions using ONLY the numbered sources provided. "
    "After every sentence that uses a source, add its number in square brackets, e.g. [1] or [2][3]. "
    "If the sources do not contain the answer, reply exactly: "
    "\"I could not find this in your documents.\" "
    "Do not use outside knowledge. Keep the answer short and in the language of the question."
)
NOT_FOUND = "I could not find this in your documents."


def _sentences(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+|\n+", text) if len(s.strip()) > 2]


def extractive_answer(question: str, hits: list[tuple[float, dict]]) -> str:
    q = set(content_tokens(question))
    best = []
    for rank, (_, c) in enumerate(hits[:3], start=1):
        for s in _sentences(c["text"]):
            st = content_tokens(s)
            overlap = len(q.intersection(st))
            if overlap:
                best.append((overlap + (0.3 if rank == 1 else 0) + overlap / max(len(st), 1), rank, s))
    if not best:
        return NOT_FOUND
    best.sort(key=lambda x: -x[0])
    top = best[0][0]
    picked, seen = [], set()
    for sc, rank, s in best[:2]:
        if sc >= 0.75 * top and s not in seen:
            picked.append(f"{s} [{rank}]")
            seen.add(s)
    return " ".join(picked)


def _post_json(url: str, payload: dict, headers: dict, timeout: int = 180) -> dict:
    req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"),
                                 headers={"Content-Type": "application/json", **headers}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8", "replace")[:300]
        raise RuntimeError(f"HTTP {e.code} from {url}: {body}")
    except urllib.error.URLError as e:
        raise RuntimeError(f"Cannot reach {url}: {e.reason}")


def llm_answer(question: str, hits, provider: str, model: str | None, base_url: str | None) -> tuple[str, dict]:
    context = "\n\n".join(f"[{i}] ({c['file']}, {c['loc']})\n{c['text']}" for i, (_, c) in enumerate(hits, start=1))
    messages = [{"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": f"Sources:\n{context}\n\nQuestion: {question}\n\n"
                 "Answer from the sources only and end each sentence with its source number, "
                 "like: <fact from a source> [1]. Never copy this format text itself."}]
    if provider == "ollama":
        host = (base_url or os.environ.get("OLLAMA_HOST") or "http://127.0.0.1:11434").rstrip("/")
        if not host.startswith("http"):
            host = "http://" + host
        model = model or os.environ.get("ASKDOCS_MODEL") or "qwen2.5:1.5b"
        r = _post_json(f"{host}/api/chat", {"model": model, "messages": messages, "stream": False,
                                            "options": {"temperature": 0}}, {})
        usage = {"prompt_tokens": r.get("prompt_eval_count"), "completion_tokens": r.get("eval_count")}
        return r["message"]["content"].strip(), {"provider": "ollama", "model": model, "endpoint": host, **usage}
    if provider == "openai":
        key = os.environ.get("OPENAI_API_KEY") or os.environ.get("ASKDOCS_API_KEY")
        base = (base_url or os.environ.get("OPENAI_BASE_URL") or "https://api.openai.com/v1").rstrip("/")
        if not key:
            raise RuntimeError("Set your API key first: OPENAI_API_KEY (see INSTALL.md, step 5).")
        model = model or os.environ.get("ASKDOCS_MODEL") or "gpt-4o-mini"
        r = _post_json(f"{base}/chat/completions", {"model": model, "messages": messages, "temperature": 0},
                       {"Authorization": f"Bearer {key}"})
        u = r.get("usage") or {}
        return (r["choices"][0]["message"]["content"] or "").strip(), {
            "provider": "openai-compatible", "model": model, "endpoint": base,
            "prompt_tokens": u.get("prompt_tokens"), "completion_tokens": u.get("completion_tokens")}
    raise RuntimeError(f"Unknown provider: {provider}")


def check_citations(answer: str, n_sources: int) -> tuple[str, list[int], list[int]]:
    """Keep only citation numbers that point to a real retrieved source."""
    cited = [int(x) for x in re.findall(r"\[(\d+)\]", answer)]
    valid = sorted({c for c in cited if 1 <= c <= n_sources})
    invalid = sorted({c for c in cited if not 1 <= c <= n_sources})
    for c in invalid:
        answer = answer.replace(f"[{c}]", "")
    return answer, valid, invalid


def _numbers(text: str) -> set[str]:
    return set(re.findall(r"\d+(?:[.,]\d+)?", text))


def attribute(answer: str, hits) -> tuple[str, list[str]]:
    """Check every sentence of an AI answer against the retrieved sources.
    - uncited sentences get the best-matching source number, if one supports them;
    - a sentence counts as supported when most of its words AND all of its numbers
      appear in the cited source(s). Unsupported sentences are returned for review."""
    if answer.strip() == NOT_FOUND:
        return answer, []
    # "fact. [1]" -> "fact [1]." so a citation stays attached to its sentence
    answer = re.sub(r"\s*([.!?])\s*((?:\[\d+\]\s*)+)", lambda m: " " + m.group(2).strip() + m.group(1) + " ", answer).strip()
    texts = [c["text"] + " " + c["title"] for _, c in hits]
    src_toks = [set(content_tokens(t)) for t in texts]
    src_nums = [_numbers(t) for t in texts]
    out, unsupported = [], []
    for sent in _sentences(answer):
        bare = re.sub(r"\[\d+\]", "", sent)
        toks = set(content_tokens(bare)) - {"ye", "yes", "no"}
        nums = _numbers(bare)
        cited = [int(x) - 1 for x in re.findall(r"\[(\d+)\]", sent) if 1 <= int(x) <= len(hits)]
        if not toks and not nums:
            out.append(sent)
            continue

        def support(idx):
            ov = len(toks & src_toks[idx]) / max(len(toks), 1)
            return ov, nums <= src_nums[idx]

        ok = False
        if cited:
            union_t = set().union(*(src_toks[i] for i in cited))
            union_n = set().union(*(src_nums[i] for i in cited))
            ok = len(toks & union_t) / max(len(toks), 1) >= 0.5 and nums <= union_n
        if not ok:  # uncited, or cited the wrong passage: point to the passage that supports it
            scores = [support(i) for i in range(len(hits))]
            best = max(range(len(hits)), key=lambda i: (scores[i][1], scores[i][0]))
            ok = scores[best][0] >= 0.5 and scores[best][1]
            if ok:
                m = re.match(r"^(.*?)\s*([.!?]*)$", bare.rstrip(), re.S)
                sent = f"{m.group(1)} [{best + 1}]{m.group(2)}"
        out.append(sent)
        if not ok:
            unsupported.append(re.sub(r"\s+([.!?])", r"\1", bare).strip())
    return " ".join(out), unsupported


def ask(index: dict, question: str, provider: str = "none", k: int = 4, model=None, base_url=None,
        bm25: BM25 | None = None) -> dict:
    bm25 = bm25 or BM25(index["chunks"])
    t0 = time.time()
    hits = bm25.search(question, k)
    if hits:  # drop weak tail matches so the source list stays relevant
        hits = [h for h in hits if h[0] >= 0.3 * hits[0][0]]
    meta = {"provider": provider}
    if not hits:
        answer = NOT_FOUND
    elif provider == "none":
        answer = extractive_answer(question, hits)
    else:
        answer, meta = llm_answer(question, hits, provider, model, base_url)
    answer, valid, invalid = check_citations(answer, len(hits))
    if hits and provider != "none":
        answer, unsupported = attribute(answer, hits)
        answer, valid, _ = check_citations(answer, len(hits))
        meta["citations_checked"] = True
        if unsupported:
            meta["unsupported_sentences"] = unsupported
    grounded = answer.strip() != NOT_FOUND
    if grounded and not valid and hits:
        meta["warning"] = "The answer has no inline citation. Check the sources below before relying on it."
    if invalid:
        meta["removed_invalid_citations"] = invalid
    sources = [{"n": i, "file": c["file"], "loc": c["loc"], "score": round(s, 2),
                "excerpt": c["text"][:300] + ("..." if len(c["text"]) > 300 else "")}
               for i, (s, c) in enumerate(hits, start=1)]
    return {"question": question, "answer": answer.strip(), "cited": valid,
            "sources": sources if grounded else [], "meta": meta, "seconds": round(time.time() - t0, 2)}


def print_result(r: dict) -> None:
    print("\n" + r["answer"])
    if r["meta"].get("warning"):
        print("\n! " + r["meta"]["warning"])
    for u in r["meta"].get("unsupported_sentences", []):
        print(f"! Not supported by the retrieved passages, check it: \"{u}\"")
    if r["sources"]:
        print("\nSources:")
        for s in r["sources"]:
            mark = "*" if s["n"] in r["cited"] else " "
            print(f" {mark}[{s['n']}] {s['file']} - {s['loc']}")
            print(f"      \"{s['excerpt'][:160]}\"")
    print()

# --------------------------------------------------------------------------- local web UI

PAGE = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Ask your docs</title><style>
body{font:16px/1.5 system-ui,sans-serif;max-width:760px;margin:24px auto;padding:0 16px;color:#111}
#log{border:1px solid #ddd;border-radius:10px;padding:12px;min-height:260px}
.m{margin:10px 0}.u{font-weight:600}.src{font-size:13px;color:#444;border-left:3px solid #3d8bff;padding-left:8px;margin:4px 0}
form{display:flex;gap:8px;margin-top:12px}input{flex:1;padding:10px;border:1px solid #bbb;border-radius:8px}
button{padding:10px 16px;border:0;border-radius:8px;background:#111;color:#fff}small{color:#666}</style></head><body>
<h1>Ask your docs</h1><p><small>Answers come only from your indexed documents. Mode: <b id="mode"></b>. Human review recommended before you rely on an answer.</small></p>
<div id="log"></div><form id="f"><input id="q" placeholder="Ask a question about your documents" autocomplete="off" required><button>Ask</button></form>
<script>
const log=document.getElementById('log'),f=document.getElementById('f'),q=document.getElementById('q');
fetch('api/info').then(r=>r.json()).then(i=>{document.getElementById('mode').textContent=i.provider+' · '+i.files+' files · '+i.chunks+' chunks'});
function add(cls,text){const d=document.createElement('div');d.className='m '+cls;d.textContent=text;log.appendChild(d);return d}
f.onsubmit=async e=>{e.preventDefault();const question=q.value.trim();if(!question)return;q.value='';add('u','You: '+question);
const w=add('a','...');try{const r=await (await fetch('api/ask',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({question})})).json();
w.textContent=r.error?('Error: '+r.error):r.answer;(r.sources||[]).forEach(s=>{const d=add('src','['+s.n+'] '+s.file+' - '+s.loc+': "'+s.excerpt.slice(0,180)+'"');w.after(d)});
if(r.meta&&r.meta.warning)add('src',r.meta.warning)}catch(err){w.textContent='Error: '+err}log.scrollTop=log.scrollHeight};
</script></body></html>"""


def serve(index: dict, provider: str, host: str, port: int, k: int, model, base_url) -> None:
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    bm25 = BM25(index["chunks"])

    class H(BaseHTTPRequestHandler):
        def _send(self, code, body, ctype="application/json"):
            data = body.encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", ctype + "; charset=utf-8")
            self.send_header("Content-Length", str(len(data)))
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            self.wfile.write(data)

        def do_GET(self):
            if self.path in ("/", "/index.html"):
                return self._send(200, PAGE, "text/html")
            if self.path == "/api/info":
                return self._send(200, json.dumps({"provider": provider, "files": len(index["files"]),
                                                   "chunks": len(index["chunks"])}))
            self._send(404, json.dumps({"error": "not found"}))

        def do_POST(self):
            if self.path != "/api/ask":
                return self._send(404, json.dumps({"error": "not found"}))
            try:
                n = min(int(self.headers.get("Content-Length", 0)), 10000)
                question = str(json.loads(self.rfile.read(n) or b"{}").get("question", ""))[:1000].strip()
                if not question:
                    return self._send(400, json.dumps({"error": "empty question"}))
                self._send(200, json.dumps(ask(index, question, provider, k, model, base_url, bm25), ensure_ascii=False))
            except Exception as e:
                self._send(500, json.dumps({"error": str(e)}))

        def log_message(self, *a):
            pass

    srv = ThreadingHTTPServer((host, port), H)
    print(f"Ask-your-docs is running on http://{host}:{port}  (mode: {provider}). Press Ctrl+C to stop.")
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")

# --------------------------------------------------------------------------- CLI


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):  # avoid crashes on Windows consoles with non-ASCII text
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass
    ap = argparse.ArgumentParser(prog="askdocs", description="Chat with your own documents, with citations.")
    ap.add_argument("--index-dir", default=os.environ.get("ASKDOCS_INDEX_DIR", DEFAULT_INDEX_DIR),
                    help="where the index is stored (default: .askdocs)")
    ap.add_argument("--version", action="version", version=f"askdocs {VERSION}")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("index", help="read and index your documents")
    p.add_argument("paths", nargs="+", help="folders or files (PDF, DOCX, MD, TXT, HTML)")

    def answer_opts(sp):
        sp.add_argument("--provider", choices=["none", "ollama", "openai"],
                        default=os.environ.get("ASKDOCS_PROVIDER", "none"))
        sp.add_argument("--model", default=None, help="model name (ollama default qwen2.5:1.5b, openai default gpt-4o-mini)")
        sp.add_argument("--base-url", default=None, help="API base URL (OpenAI-compatible) or Ollama host")
        sp.add_argument("-k", type=int, default=4, help="number of passages to retrieve (default 4)")

    p = sub.add_parser("ask", help="ask one question")
    p.add_argument("question")
    p.add_argument("--json", action="store_true", help="print machine-readable JSON")
    answer_opts(p)
    p = sub.add_parser("chat", help="interactive chat in the terminal")
    answer_opts(p)
    p = sub.add_parser("serve", help="local web chat in your browser")
    p.add_argument("--host", default="127.0.0.1")
    p.add_argument("--port", type=int, default=8765)
    answer_opts(p)
    p = sub.add_parser("export-web", help="write a docs.json for the static browser widget")
    p.add_argument("out", help="output path, e.g. ../data/docs.json")
    sub.add_parser("info", help="show what is indexed")

    a = ap.parse_args(argv)
    if a.cmd == "index":
        d = build_index(a.paths, a.index_dir)
        print(f"Indexed {len(d['files'])} file(s) into {len(d['chunks'])} passages -> {Path(a.index_dir) / 'index.json'}")
        for f in d["files"]:
            print(f"  + {f}")
        for s in d["skipped"]:
            print(f"  ! skipped {s['file']}: {s['reason']}")
        return 0
    index = load_index(a.index_dir)
    if a.cmd == "info":
        print(json.dumps({"files": index["files"], "skipped": index["skipped"], "chunks": len(index["chunks"]),
                          "created": index["created"]}, indent=1))
        return 0
    if a.cmd == "export-web":
        docs = [{"id": c["id"], "title": c["title"], "source": f"{c['file']} - {c['loc']}", "text": c["text"]}
                for c in index["chunks"]]
        Path(a.out).parent.mkdir(parents=True, exist_ok=True)
        Path(a.out).write_text(json.dumps(docs, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"Wrote {len(docs)} passages to {a.out}")
        return 0
    if a.cmd == "ask":
        try:
            r = ask(index, a.question, a.provider, a.k, a.model, a.base_url)
        except RuntimeError as e:
            print(f"Error: {e}", file=sys.stderr)
            return 2
        print(json.dumps(r, ensure_ascii=False, indent=1)) if a.json else print_result(r)
        return 0
    if a.cmd == "chat":
        bm25 = BM25(index["chunks"])
        print(f"Ask-your-docs chat - {len(index['files'])} files, mode: {a.provider}. Type 'exit' to quit.")
        while True:
            try:
                q = input("You> ").strip()
            except (EOFError, KeyboardInterrupt):
                print()
                return 0
            if q.lower() in ("exit", "quit", ":q"):
                return 0
            if q:
                try:
                    print_result(ask(index, q, a.provider, a.k, a.model, a.base_url, bm25))
                except RuntimeError as e:
                    print(f"Error: {e}")
    if a.cmd == "serve":
        serve(index, a.provider, a.host, a.port, a.k, a.model, a.base_url)
        return 0
    return 1


if __name__ == "__main__":
    sys.exit(main())
